from pathlib import Path
from tempfile import TemporaryDirectory
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw
from copy import deepcopy

ROOT=Path(__file__).resolve().parents[1]
DEST=ROOT/'dist'/'downloads'
DEST.mkdir(parents=True,exist_ok=True)
asset_workspace=TemporaryDirectory(prefix='vagas-ti-curriculos-')
ASSETS=Path(asset_workspace.name)
ASSETS.mkdir(parents=True,exist_ok=True)

def icon(name):
    path=ASSETS/(name+'.png')
    im=Image.new('RGBA',(120,120),(255,255,255,0));d=ImageDraw.Draw(im)
    ink='#111820';w=5
    def line(points): d.line(points,fill=ink,width=w,joint='curve')
    if name=='goal':
        d.ellipse((18,18,102,102),outline=ink,width=w);d.ellipse((38,38,82,82),outline=ink,width=w)
        line([(60,60),(100,20)]);line([(80,20),(100,20),(100,40)])
    elif name=='education':
        line([(10,44),(60,20),(110,44),(60,69),(10,44)])
        line([(30,56),(30,83),(60,97),(90,83),(90,56)]);line([(110,44),(110,85)])
    elif name=='code':
        line([(40,30),(13,60),(40,90)]);line([(80,30),(107,60),(80,90)]);line([(70,20),(50,100)])
    elif name=='project':
        d.rounded_rectangle((14,20,106,90),radius=7,outline=ink,width=w)
        line([(14,41),(106,41)]);line([(42,56),(32,66),(42,77)]);line([(78,56),(88,66),(78,77)])
        line([(46,105),(74,105)]);line([(60,91),(60,105)])
    elif name=='course':
        line([(16,22),(53,22),(60,30),(67,22),(104,22),(104,96),(68,96),(60,103),(52,96),(16,96),(16,22)])
        line([(60,30),(60,103)]);line([(28,42),(45,42)]);line([(75,42),(92,42)])
    elif name=='language':
        d.ellipse((15,15,105,105),outline=ink,width=w);d.ellipse((39,15,81,105),outline=ink,width=w)
        line([(15,60),(105,60)]);line([(25,35),(95,35)]);line([(25,85),(95,85)])
    elif name=='clock':
        d.ellipse((15,15,105,105),outline=ink,width=w);line([(60,32),(60,62),(82,75)])
    elif name=='person':
        d.ellipse((42,12,78,48),outline=ink,width=w);d.arc((20,58,100,128),180,360,fill=ink,width=w)
        line([(20,91),(20,105)]);line([(100,91),(100,105)])
    elif name=='work':
        d.rounded_rectangle((12,38,108,103),radius=7,outline=ink,width=w)
        line([(40,38),(40,20),(80,20),(80,38)]);line([(12,62),(108,62)])
        d.rectangle((51,55,69,72),fill='white',outline=ink,width=w)
    im.save(path);return path

def background(doc,name):
    colors=('#D8E9ED','#F7D6D4','#B2A07F') if name.endswith('estagio') else ('#DCE4F1','#D6EADF','#42556E')
    im=Image.new('RGBA',(1700,2200),(255,255,255,0));d=ImageDraw.Draw(im)
    d.rectangle((0,0,45,840),fill=colors[0]);d.rectangle((1090,0,1600,28),fill=colors[1]);d.rectangle((1120,2170,1700,2200),fill=colors[2])
    path=ASSETS/(name+'-background.png');im.save(path)
    sec=doc.sections[0];p=sec.header.paragraphs[0]
    inline=p.add_run().add_picture(str(path),width=sec.page_width,height=sec.page_height)._inline
    anchor=OxmlElement('wp:anchor')
    for k,v in {'distT':'0','distB':'0','distL':'0','distR':'0','simplePos':'0','relativeHeight':'0','behindDoc':'1','locked':'0','layoutInCell':'1','allowOverlap':'1'}.items(): anchor.set(k,v)
    sp=OxmlElement('wp:simplePos');sp.set('x','0');sp.set('y','0');anchor.append(sp)
    for axis in ['H','V']:
        pos=OxmlElement('wp:position'+axis);pos.set('relativeFrom','page');off=OxmlElement('wp:posOffset');off.text='0';pos.append(off);anchor.append(pos)
    anchor.append(deepcopy(inline.find(qn('wp:extent'))));anchor.append(OxmlElement('wp:wrapNone'))
    for tag in ['wp:docPr','wp:cNvGraphicFramePr','a:graphic']: anchor.append(deepcopy(inline.find(qn(tag))))
    inline.getparent().replace(inline,anchor)

ICON_MAP={'Objetivo':'goal','Resumo':'person','Formação acadêmica':'education','Conhecimentos técnicos':'code','Projetos e atividades acadêmicas':'project','Projetos selecionados':'project','Cursos complementares':'course','Idiomas':'language','Disponibilidade':'clock','Experiência e atividades':'work','Conhecimentos e cursos':'code','Idiomas e disponibilidade':'language'}

def build(name, sections):
    doc=Document()
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.58);sec.bottom_margin=Inches(.6)
    sec.left_margin=sec.right_margin=Inches(.6)
    background(doc,name)
    for stylename in ['Normal','Title','Subtitle','Heading 1','Heading 2','List Bullet']:
        style=doc.styles[stylename]
        style.font.name='Calibri';style.font.color.rgb=RGBColor.from_string('333B43')
        style.font.size=Pt(11)
        style.paragraph_format.space_after=Pt(4)
        style.paragraph_format.line_spacing=1.05
    doc.styles['Title'].font.size=Pt(30)
    doc.styles['Title'].font.color.rgb=RGBColor(0,0,0)
    doc.styles['Title'].font.bold=True
    doc.styles['Title'].paragraph_format.space_after=Pt(6)
    doc.styles['Heading 1'].font.size=Pt(12)
    doc.styles['Heading 1'].font.color.rgb=RGBColor(0,0,0)
    doc.styles['Heading 1'].font.bold=True
    doc.styles['Heading 1'].paragraph_format.space_before=Pt(13)
    doc.styles['Heading 1'].paragraph_format.space_after=Pt(7)
    doc.styles['Heading 1'].paragraph_format.keep_with_next=True
    doc.add_paragraph('Seu nome completo',style='Title')
    p=doc.add_paragraph();r=p.add_run('ESTUDANTE DE [CURSO]  |  ESTÁGIO EM [ÁREA]');r.bold=True;r.font.size=Pt(10);r.font.color.rgb=RGBColor.from_string('315C65' if name.endswith('estagio') else '42556E')
    for text in ['[Cidade e UF]  •  [Telefone com DDD]  •  [E-mail profissional]','[LinkedIn, se tiver]  •  [GitHub ou portfólio, se tiver]']:
        p=doc.add_paragraph(text);p.paragraph_format.space_after=Pt(3)
        for r in p.runs:r.font.size=Pt(10)
    for heading,paragraphs in sections:
        p=doc.add_paragraph(style='Heading 1');p.add_run().add_picture(str(icon(ICON_MAP[heading])),width=Inches(.18));p.add_run('  '+heading.upper())
        borders=OxmlElement('w:pBdr');bottom=OxmlElement('w:bottom')
        for k,v in {'val':'single','sz':'6','space':'4','color':'202831'}.items():bottom.set(qn('w:'+k),v)
        borders.append(bottom);p._p.get_or_add_pPr().append(borders)
        for line in paragraphs:
            if line.startswith('* '):
                p=doc.add_paragraph(line[2:],style='List Bullet')
            else:
                p=doc.add_paragraph(line)
                p.paragraph_format.left_indent=Inches(.26)
                if line.startswith(('[Nome do projeto','[Nome da atividade','[Nome do curso]','[Curso]','[Função')):
                    for r in p.runs:r.bold=True;r.font.color.rgb=RGBColor.from_string('18212A')
            p.paragraph_format.widow_control=True
    doc.core_properties.title='Modelo de currículo para estágio em TI'
    doc.core_properties.subject=name
    doc.core_properties.author='Vagas TI'
    doc.core_properties.comments='Modelo original. Substitua os campos entre colchetes por informações verdadeiras e remova seções que não se aplicam antes de enviar.'
    for style in doc.styles:
        if style.type==1:
            lang=OxmlElement('w:lang');lang.set(qn('w:val'),'pt-BR')
            style.element.get_or_add_rPr().append(lang)
    target=DEST/(name+'.docx')
    doc.save(target)
    (DEST/(name+'.txt')).write_text('\n\n'.join(p.text for p in doc.paragraphs)+'\n',encoding='utf-8')
    print(target)

build('curriculo-primeiro-estagio',[
('Objetivo',['Estágio em [área ou cargo da vaga].']),
('Formação acadêmica',['[Nome do curso] — [Instituição]','[Semestre ou período atual] | Conclusão prevista: [mês e ano]']),
('Conhecimentos técnicos',['[Linguagens e ferramentas que você já utilizou]','[Descreva onde aplicou esses conhecimentos: disciplina, exercício ou projeto.]']),
('Projetos e atividades acadêmicas',['[Nome da atividade ou projeto] — [disciplina, curso ou iniciativa] | [ano]','* [O que você desenvolveu ou executou, usando quais ferramentas.]','* [Sua contribuição e o resultado que consegue demonstrar, sem inventar números.]','[Link público do projeto, se houver]']),
('Cursos complementares',['[Nome do curso] — [Instituição] | [carga horária] | [ano]','[Outro curso relevante, se houver]']),
('Idiomas',['[Idioma] — [nível que consegue demonstrar e contexto de uso]']),
('Disponibilidade',['[Dias e horários compatíveis com as aulas] | [Presencial, híbrido ou remoto]']),
])
build('curriculo-com-projetos',[
('Objetivo',['Estágio em [desenvolvimento, dados, suporte ou outra área da vaga].']),
('Resumo',['Estudante de [curso], com prática em [tecnologias] em projetos [acadêmicos ou pessoais]. Interesse em [área relacionada à vaga].']),
('Projetos selecionados',['[Nome do projeto 1] | [período] | [link do repositório ou demonstração]','* [Problema que o projeto resolve e sua contribuição individual.]','* [Tecnologias usadas, testes realizados e resultado verificável.]','[Nome do projeto 2] | [período] | [link do repositório ou demonstração]','* [O que você implementou e como organizou ou validou a solução.]']),
('Experiência e atividades',['[Função em emprego, monitoria, voluntariado ou extensão] — [Organização]','[Mês/ano de início] a [mês/ano de fim ou atual]','* [Entrega, atendimento ou colaboração relevante para a oportunidade.]']),
('Formação acadêmica',['[Curso] — [Instituição] | [período atual] | Conclusão prevista: [mês/ano]']),
('Conhecimentos e cursos',['[Linguagens, ferramentas e práticas utilizadas nos projetos]','[Curso relevante] — [Instituição] | [ano]']),
('Idiomas e disponibilidade',['[Idioma e nível] | [Dias e horários disponíveis] | [Modelo de trabalho]']),
])

